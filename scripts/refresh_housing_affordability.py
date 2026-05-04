"""
Refresh data/static/housing_affordability.json from ONS source.

ONS publishes the House Price to Workplace-Based Earnings Ratio annually
(typically late March, year ending the previous September). Run this script
when a new year is released — there's no live cadence to schedule.

Source page:
  https://www.ons.gov.uk/peoplepopulationandcommunity/housing/datasets/
    ratioofhousepricetoworkplacebasedearningslowerquartileandmedian

The "current" workbook URL is stable across releases (ONS overwrites it
when a new year drops):

  https://www.ons.gov.uk/file?uri=/peoplepopulationandcommunity/housing/datasets/
    ratioofhousepricetoworkplacebasedearningslowerquartileandmedian/current/
    aff1ratioofhousepricetoworkplacebasedearnings.xlsx

Output JSON shape mirrors data/static/imd_borough.json so the API router
(api/routers/boroughs.py:_load_housing_data) can read both with the same idiom.

Usage:
  python -m scripts.refresh_housing_affordability
"""

from __future__ import annotations

import json
import logging
import sys
from pathlib import Path
from urllib.request import urlretrieve

import openpyxl

logger = logging.getLogger(__name__)

ONS_URL = (
    "https://www.ons.gov.uk/file?uri="
    "/peoplepopulationandcommunity/housing/datasets/"
    "ratioofhousepricetoworkplacebasedearningslowerquartileandmedian/"
    "current/aff1ratioofhousepricetoworkplacebasedearnings.xlsx"
)

# Sheet "5c" = median affordability ratio by Local Authority District.
SHEET = "5c"

# London borough ONS codes (32 boroughs + City of London = 33).
LONDON_CODES = (
    {f"E0900000{i}" for i in range(1, 10)}
    | {f"E090000{i}" for i in range(10, 34)}
)

REPO_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = REPO_ROOT / "data" / "static" / "housing_affordability.json"


def _download(dest: Path) -> None:
    logger.info("Downloading %s", ONS_URL)
    urlretrieve(ONS_URL, dest)


def _extract_london_rows(xlsx_path: Path) -> tuple[list[tuple], int]:
    """Extract (code, name, latest_ratio, 5yr_avg) for the 33 London boroughs.

    Returns the rows and the latest year covered (e.g. 2025).
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[SHEET]

    # Row 2 holds the header. Column 2 = LA code, column 3 = LA name. The
    # last numeric-string column header is the latest year; the column after
    # is "5-Year Average".
    header = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
    year_cols = [(i, int(v)) for i, v in enumerate(header) if isinstance(v, str) and v.isdigit()]
    if not year_cols:
        raise RuntimeError("No year columns found in header — ONS workbook layout changed")
    latest_idx, latest_year = year_cols[-1]
    five_yr_idx = next(
        (i for i, v in enumerate(header) if v == "5-Year Average"),
        None,
    )

    rows: list[tuple] = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        code = row[2]
        name = row[3]
        if code in LONDON_CODES:
            ratio = row[latest_idx]
            five_yr = row[five_yr_idx] if five_yr_idx is not None else None
            rows.append((code, name, ratio, five_yr))

    if len(rows) != 33:
        raise RuntimeError(f"Expected 33 London boroughs, found {len(rows)}")
    return rows, latest_year


def build_json(rows: list[tuple], year: int) -> dict:
    return {
        "_metadata": {
            "source": "ONS House Price to Workplace-Based Earnings Ratio (Table 5c, median)",
            "url": (
                "https://www.ons.gov.uk/peoplepopulationandcommunity/housing/datasets/"
                "ratioofhousepricetoworkplacebasedearningslowerquartileandmedian"
            ),
            "year": year,
            "year_period": f"Year ending September {year}",
            "definition": (
                "Ratio of median house price to median gross annual "
                "workplace-based earnings"
            ),
            "geography": "Local authority district (32 London boroughs + City of London)",
            "table": "5c",
        },
        "boroughs": {
            code: {
                "name": name,
                f"ratio_{year}": float(ratio) if isinstance(ratio, (int, float)) else None,
                "ratio_5yr_avg": float(five_yr) if isinstance(five_yr, (int, float)) else None,
            }
            for code, name, ratio, five_yr in rows
        },
    }


def main() -> int:
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    tmp_xlsx = Path("/tmp/_ons_affordability.xlsx")
    _download(tmp_xlsx)
    rows, year = _extract_london_rows(tmp_xlsx)
    payload = build_json(rows, year)
    OUTPUT_PATH.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    logger.info("Wrote %s (%d boroughs, year=%d)", OUTPUT_PATH, len(rows), year)
    return 0


if __name__ == "__main__":
    sys.exit(main())
