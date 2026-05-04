"""
Refresh data/static/ons_health_index.json from ONS source.

ONS publishes the Health Index for England as a workbook with multiple
sheets covering the index from 2015 onwards. The latest published year as
of 2026-05 is 2021 (released June 2023). The Health Index is updated
infrequently — re-run this script when ONS publishes a new vintage.

Source page:
  https://www.ons.gov.uk/peoplepopulationandcommunity/
    healthandsocialcare/healthandwellbeing/datasets/healthindexscoresengland

The "current" workbook URL is stable across releases:

  https://www.ons.gov.uk/file?uri=/peoplepopulationandcommunity/
    healthandsocialcare/healthandwellbeing/datasets/healthindexscoresengland/
    current/healthindexscoresengland.xlsx

Sheet "Table_2_Index_scores" holds the overall Health Index score per LA
across years 2015–2021. Score of 100 = England-2015 baseline.

Output JSON shape mirrors data/static/imd_borough.json so the API router
(api/routers/boroughs.py:_load_health_index_data) can read it with the
same idiom.

Note: City of London (E09000001) is NOT included in the Health Index
dataset — ONS excludes it because of its tiny resident population and
overwhelming commuter workforce. Output therefore has 32 boroughs, not 33.

Usage:
  python -m scripts.refresh_ons_health_index
"""

from __future__ import annotations

import json
import logging
import sys
from pathlib import Path
from urllib.request import urlretrieve

import openpyxl

logger = logging.getLogger(__name__)

ONS_URL = (
    "https://www.ons.gov.uk/file?uri="
    "/peoplepopulationandcommunity/healthandsocialcare/healthandwellbeing/"
    "datasets/healthindexscoresengland/current/healthindexscoresengland.xlsx"
)

SHEET = "Table_2_Index_scores"

# London borough ONS codes (32 borough councils + City of London = 33;
# CoL is excluded from the Health Index — see module docstring).
LONDON_CODES = (
    {f"E0900000{i}" for i in range(1, 10)}
    | {f"E090000{i}" for i in range(10, 34)}
)

REPO_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = REPO_ROOT / "data" / "static" / "ons_health_index.json"


def _download(dest: Path) -> None:
    logger.info("Downloading %s", ONS_URL)
    urlretrieve(ONS_URL, dest)


def _extract_london_rows(xlsx_path: Path) -> tuple[list[dict], int]:
    """Extract all years of Health Index for London boroughs.

    Returns (rows, latest_year).
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[SHEET]

    # Header on row 3: Area Code, Area Name, Area Type, 2015..2021
    header = next(ws.iter_rows(min_row=3, max_row=3, values_only=True))
    year_cols: list[tuple[int, int]] = []
    for i, v in enumerate(header):
        if isinstance(v, int) and 2000 < v < 2100:
            year_cols.append((i, v))
        elif isinstance(v, str) and v.isdigit() and 2000 < int(v) < 2100:
            year_cols.append((i, int(v)))
    if not year_cols:
        raise RuntimeError("No year columns found — workbook layout changed")
    latest_year = year_cols[-1][1]

    rows: list[dict] = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        code = row[0]
        if code in LONDON_CODES:
            scores = {y: row[i] for i, y in year_cols}
            rows.append({
                "code": code,
                "name": row[1],
                "type": row[2],
                "scores": scores,
            })

    return rows, latest_year


def build_json(rows: list[dict], latest_year: int) -> dict:
    return {
        "_metadata": {
            "source": "ONS Health Index for England (Table 2 — Overall Index scores)",
            "url": (
                "https://www.ons.gov.uk/peoplepopulationandcommunity/"
                "healthandsocialcare/healthandwellbeing/datasets/healthindexscoresengland"
            ),
            "year": latest_year,
            "released": "2023-06-16",
            "definition": (
                "Composite index of population health based on independent inputs: "
                "life expectancy, healthy life expectancy, mortality from specific "
                "causes, mental health prevalence, behaviours (smoking, drinking, "
                "physical activity), and broader health determinants. England 2015 = 100."
            ),
            "geography": "Lower-Tier Local Authority (LTLA). City of London omitted by ONS.",
            "table": "Table_2_Index_scores",
            "lag_note": (
                f"Latest published year is {latest_year}. The Health Index is "
                "published with a multi-year delay — this is the freshest available "
                "vintage as of refresh time."
            ),
        },
        "boroughs": {
            r["code"]: {
                "name": r["name"],
                f"score_{latest_year}": (
                    float(r["scores"][latest_year])
                    if isinstance(r["scores"][latest_year], (int, float)) else None
                ),
                "score_2020": (
                    float(r["scores"][2020])
                    if r["scores"].get(2020) is not None
                    and isinstance(r["scores"][2020], (int, float)) else None
                ),
                "score_2015": (
                    float(r["scores"][2015])
                    if r["scores"].get(2015) is not None
                    and isinstance(r["scores"][2015], (int, float)) else None
                ),
            }
            for r in rows
        },
    }


def main() -> int:
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    tmp_xlsx = Path("/tmp/_ons_health_index.xlsx")
    _download(tmp_xlsx)
    rows, latest = _extract_london_rows(tmp_xlsx)
    payload = build_json(rows, latest)
    OUTPUT_PATH.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    logger.info("Wrote %s (%d boroughs, latest year=%d)", OUTPUT_PATH, len(rows), latest)
    return 0


if __name__ == "__main__":
    sys.exit(main())
