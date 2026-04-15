"""
Fetch English Indices of Deprivation 2025 for London boroughs.

Downloads the Local Authority District summaries XLSX from GOV.UK,
parses the IMD sheet and all 7 domain sheets for the 33 London boroughs,
and writes a static JSON file. Used by the deprivation agent tool.

Data source:
    MHCLG English Indices of Deprivation 2025 (File 10 — LAD summaries).
    Free, no API key required.

Run once (data is static — published roughly every 6 years):
    python scripts/fetch_imd.py

Safe to re-run — overwrites data/static/imd_borough.json.
"""

import json
import logging
import sys
import tempfile
import time
from pathlib import Path

import openpyxl
import requests

logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s: %(message)s")
logger = logging.getLogger(__name__)

PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUTPUT_PATH = PROJECT_ROOT / "data" / "static" / "imd_borough.json"

# GOV.UK — File 10: Local Authority District summaries (lower tier) v2
XLSX_URL = (
    "https://assets.publishing.service.gov.uk/media/"
    "6917412ebc34c86ce4e6e7fc/"
    "File_10_-_IoD2025_Local_Authority_District_Summaries__lower-tier__v2.xlsx"
)

DATASET_URL = "https://www.gov.uk/government/statistics/english-indices-of-deprivation-2025"

# All 33 London borough ONS codes
EXPECTED_CODES = {f"E0900{str(i).zfill(4)}" for i in range(1, 34)}

# Domain sheets and their column mappings
# Each domain sheet has: code[0], name[1], avg_rank[2], rank_of_avg_rank[3], avg_score[4], rank_of_avg_score[5]
DOMAIN_SHEETS = {
    "Income": "income",
    "Employment": "employment",
    "Education": "education",
    "Health": "health",
    "Crime": "crime",
    "Barriers": "barriers",
    "Living": "living",
}


def download_xlsx(url: str) -> Path:
    """Download the XLSX to a temporary file and return its path."""
    logger.info("Downloading IMD 2025 XLSX from GOV.UK...")
    resp = None
    for attempt in range(3):
        try:
            resp = requests.get(url, timeout=60)
            resp.raise_for_status()
            break
        except requests.RequestException as e:
            if attempt == 2:
                logger.error("Download failed after 3 attempts: %s", e)
                raise
            logger.warning("Download attempt %d failed: %s — retrying", attempt + 1, e)
            time.sleep(2 ** attempt)

    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.write(resp.content)
    tmp.close()
    logger.info("Downloaded %.1f KB to %s", len(resp.content) / 1024, tmp.name)
    return Path(tmp.name)


def parse_imd_sheet(wb: openpyxl.Workbook) -> dict[str, dict]:
    """Parse the IMD sheet for London boroughs.

    Returns dict keyed by ONS code with imd_rank, imd_score, pct_most_deprived_10.
    """
    ws = wb["IMD"]
    boroughs: dict[str, dict] = {}

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue  # header

        code = row[0]
        if not code or not str(code).startswith("E09"):
            continue

        code = str(code).strip()
        boroughs[code] = {
            "name": str(row[1]).strip(),
            "imd_rank": int(row[3]) if row[3] is not None else None,  # rank of average rank
            "imd_score": round(float(row[4]), 2) if row[4] is not None else None,  # average score
            "pct_most_deprived_10": round(float(row[6]) * 100, 1) if row[6] is not None else None,
        }

    logger.info("IMD sheet: %d London boroughs", len(boroughs))
    return boroughs


def parse_domain_sheet(wb: openpyxl.Workbook, sheet_name: str) -> dict[str, dict]:
    """Parse a domain sheet for London boroughs.

    Returns dict keyed by ONS code with rank and score for this domain.
    """
    ws = wb[sheet_name]
    results: dict[str, dict] = {}

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue

        code = row[0]
        if not code or not str(code).startswith("E09"):
            continue

        code = str(code).strip()
        results[code] = {
            "rank": int(row[3]) if row[3] is not None else None,  # rank of average rank
            "score": round(float(row[4]), 3) if row[4] is not None else None,  # average score
        }

    return results


def main() -> None:
    """Download XLSX, parse all sheets, validate, and write JSON."""
    # Download or use local copy
    xlsx_path = None
    try:
        xlsx_path = download_xlsx(XLSX_URL)
    except Exception:
        local = Path("/tmp/imd_2025_lad.xlsx")
        if local.exists():
            logger.info("Using local copy at %s", local)
            xlsx_path = local
        else:
            logger.error("No XLSX available — cannot proceed")
            sys.exit(1)

    wb = openpyxl.load_workbook(xlsx_path, read_only=True)

    # Parse IMD headline data
    boroughs = parse_imd_sheet(wb)

    # Parse all 7 domain sheets
    for sheet_name, domain_key in DOMAIN_SHEETS.items():
        if sheet_name not in wb.sheetnames:
            logger.warning("Sheet '%s' not found — skipping domain", sheet_name)
            continue

        domain_data = parse_domain_sheet(wb, sheet_name)
        for code, domain_vals in domain_data.items():
            if code in boroughs:
                if "domains" not in boroughs[code]:
                    boroughs[code]["domains"] = {}
                boroughs[code]["domains"][domain_key] = domain_vals

        logger.info("Domain %s: %d London boroughs", domain_key, len(domain_data))

    wb.close()

    # Sanity checks
    found_codes = set(boroughs.keys())
    missing = EXPECTED_CODES - found_codes
    if missing:
        logger.error("Missing %d borough(s): %s", len(missing), sorted(missing))
        sys.exit(1)

    if len(boroughs) != 33:
        logger.error("Expected 33 boroughs, got %d", len(boroughs))
        sys.exit(1)

    # Check all boroughs have 7 domains
    for code, b in boroughs.items():
        domains = b.get("domains", {})
        if len(domains) != 7:
            logger.warning(
                "%s (%s): only %d/7 domains parsed",
                b["name"], code, len(domains),
            )

    # Log summary
    ranks = [b["imd_rank"] for b in boroughs.values() if b["imd_rank"]]
    most_deprived = min(boroughs.values(), key=lambda b: b["imd_rank"] or 999)
    least_deprived = max(boroughs.values(), key=lambda b: b["imd_rank"] or 0)
    logger.info(
        "IMD rank range: %d (%s) – %d (%s) out of 296 LADs nationally",
        most_deprived["imd_rank"], most_deprived["name"],
        least_deprived["imd_rank"], least_deprived["name"],
    )

    # Write output
    output = {
        "_metadata": {
            "source": "English Indices of Deprivation 2025 (MHCLG)",
            "url": DATASET_URL,
            "published": "October 2025",
            "total_lads": 296,
            "domains": list(DOMAIN_SHEETS.values()),
        },
        "boroughs": boroughs,
    }

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(output, f, indent=2, ensure_ascii=False)

    # Clean up temp file
    if xlsx_path != Path("/tmp/imd_2025_lad.xlsx"):
        xlsx_path.unlink(missing_ok=True)

    logger.info("Wrote %s (%d boroughs, 7 domains each)", OUTPUT_PATH, len(boroughs))


if __name__ == "__main__":
    sys.exit(main() or 0)
